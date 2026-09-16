"""B2B sourcing: dedup, criteria filtering, and fail-closed source handling."""

from __future__ import annotations

import sys
from datetime import datetime, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from fexlead.b2b.pipeline import run_scrape_job  # noqa: E402
from fexlead.b2b.schema import BizContact, Company, SourceType  # noqa: E402
from fexlead.b2b.sources import InMemorySource, ScrapeCriteria, run_sources  # noqa: E402

NOW = datetime(2026, 9, 16, tzinfo=timezone.utc)


def mk(name, state, **kw):
    return Company(name=name, source=SourceType.MOCK, state=state, **kw)


def sample():
    return [
        mk("Acme Robotics", "TX", domain="acme.io", industry="Software", phone="5125550100",
           signal="Raised $4M", signal_date=NOW,
           contacts=[BizContact(full_name="Dana Lee", title="CEO", email="dana@acme.io", provenance="public_filing")]),
        mk("Bayou Health", "LA", domain="bayou.health", industry="Healthcare", phone="5045550100"),
        mk("Acme Robotics", "TX", domain="acme.io"),  # duplicate by domain
        mk("Lone Star SaaS", "TX", domain="lonestar.app", industry="Software"),  # no contact channel
    ]


class ExplodingSource:
    name = "flaky_api"

    def fetch(self, criteria):
        raise ConnectionError("endpoint 500")


def test_dedup_by_domain():
    companies, report = run_sources([InMemorySource(companies=sample())], ScrapeCriteria(count=50))
    names = [c.name for c in companies]
    assert names.count("Acme Robotics") == 1
    assert report.fetched == 4
    assert report.after_dedup == 3


def test_state_filter():
    companies, _ = run_sources([InMemorySource(companies=sample())], ScrapeCriteria(states=["TX"]))
    assert {c.state for c in companies} == {"TX"}


def test_industry_filter():
    companies, _ = run_sources([InMemorySource(companies=sample())], ScrapeCriteria(industries=["software"]))
    assert all("Software" == c.industry for c in companies)


def test_require_contact_channel_drops_unreachable():
    companies, _ = run_sources([InMemorySource(companies=sample())],
                               ScrapeCriteria(states=["TX"], require_contact_channel=True))
    names = {c.name for c in companies}
    assert "Acme Robotics" in names       # has phone + contact email
    assert "Lone Star SaaS" not in names   # no phone, no contact


def test_failed_source_is_reported_not_hidden():
    """A source that raises must be recorded, never silently degraded to empty."""
    companies, report = run_sources(
        [InMemorySource(companies=sample()), ExplodingSource()], ScrapeCriteria(count=50))
    assert len(report.source_failures) == 1
    assert "flaky_api" in report.source_failures[0]
    assert len(companies) == 3  # the working source still delivered


def test_count_cap():
    companies, _ = run_sources([InMemorySource(companies=sample())], ScrapeCriteria(count=2))
    assert len(companies) <= 2


def test_scrape_job_writes_excel(tmp_path):
    out = tmp_path / "leads.xlsx"
    result = run_scrape_job([InMemorySource(companies=sample())],
                            ScrapeCriteria(states=["TX"], require_contact_channel=True),
                            output_path=str(out), now=NOW)
    assert out.exists()
    assert result.output_path == str(out)
    from openpyxl import load_workbook
    ws = load_workbook(out)["Leads"]
    assert ws.max_row >= 2
    hdr = [c.value for c in ws[1]]
    assert "Data Provenance" in hdr  # buyers must see where personal data came from
