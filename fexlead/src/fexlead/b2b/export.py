"""Clean Excel export for a B2B scrape job -- what the dashboard's 'Download' button produces."""

from __future__ import annotations

from datetime import datetime, timezone
from typing import Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..verification import verify_email, verify_phone
from .schema import Company

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)

COLUMNS = [
    ("Company", 26), ("Industry", 22), ("Signal", 26), ("Signal Date", 13),
    ("Phone", 15), ("Phone Check", 12), ("Website", 22),
    ("Address", 24), ("City", 16), ("State", 7), ("ZIP", 8),
    ("Contact Name", 20), ("Title", 18), ("Contact Email", 26), ("Email Check", 12),
    ("Data Provenance", 18), ("Source", 16), ("Source URL", 30),
]


def build_workbook(companies: Sequence[Company], now: Optional[datetime] = None,
                   check_email_deliverability: bool = False) -> Workbook:
    now = now or datetime.now(timezone.utc)
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    for i, (title, width) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=title)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"

    for company in companies:
        contact = company.contacts[0] if company.contacts else None
        phone = company.phone or (contact.phone if contact else None)
        phone_check = verify_phone(phone).status.value if phone else "no phone"
        email = contact.email if contact else None
        email_check = verify_email(email, check_deliverability=check_email_deliverability).status.value if email else "no email"
        ws.append([
            company.name, company.industry, company.signal,
            company.signal_date.strftime("%Y-%m-%d") if company.signal_date else None,
            phone, phone_check, company.domain,
            company.address1, company.city, company.state, company.postal_code,
            contact.full_name if contact else None,
            contact.title if contact else None,
            email, email_check,
            contact.provenance if contact else None,
            company.source.value, company.source_url,
        ])
    return wb


def write_xlsx(companies: Sequence[Company], path: str, **kw) -> str:
    build_workbook(companies, **kw).save(path)
    return path
